from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import from_excel
from pydantic import ValidationError

from schemas import CURRENCIES, TransactionCreate


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 2000
MAX_COLUMNS = 32
MAX_UNCOMPRESSED_XLSX_BYTES = 25 * 1024 * 1024
MAX_XLSX_ENTRIES = 200

FIELDS = (
    "transaction_date",
    "kind",
    "amount",
    "currency",
    "category",
    "merchant",
    "note",
    "tags",
)
REQUIRED_FIELDS = ("transaction_date", "kind", "amount", "category")
FIELD_LABELS = {
    "transaction_date": "Transaction Date",
    "kind": "Type",
    "amount": "Amount",
    "currency": "Currency",
    "category": "Category",
    "merchant": "Merchant",
    "note": "Note / Description",
    "tags": "Tags",
}
HEADER_ALIASES = {
    "transaction_date": {"date", "transactiondate", "txdate"},
    "kind": {"type", "kind", "transactiontype"},
    "amount": {"amount", "value"},
    "currency": {"currency", "curr"},
    "category": {"category", "typecategory"},
    "merchant": {"merchant", "payee", "vendor"},
    "note": {"note", "description", "memo", "details"},
    "tags": {"tag", "tags"},
}
TEMPLATE_HEADERS = [FIELD_LABELS[field] for field in FIELDS]


class ImportFileError(ValueError):
    """Raised when an upload is not a supported transaction file."""


def normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


def _string_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        raise ValueError("Use an Excel date cell or a supported date string, not a plain number")
    if not isinstance(value, str) or not value.strip():
        raise ValueError("Date is required")

    text = value.strip()
    try:
        return date.fromisoformat(text)
    except ValueError:
        pass

    for date_format in ("%m/%d/%Y", "%m/%d/%y", "%d-%b-%Y", "%d %b %Y"):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    raise ValueError("Use YYYY-MM-DD, MM/DD/YYYY, or DD-Mon-YYYY for dates")


def _parse_amount(value: Any) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise ValueError("Amount must be a number greater than zero")
    try:
        amount = Decimal(str(value).strip())
    except (InvalidOperation, ValueError, AttributeError):
        raise ValueError("Amount must be a valid number") from None
    if not amount.is_finite() or amount <= 0:
        raise ValueError("Amount must be a finite number greater than zero")
    if amount >= Decimal("10000000000000000"):
        raise ValueError("Amount is larger than the supported limit")
    if amount.as_tuple().exponent < -2:
        raise ValueError("Amount cannot have more than two decimal places")
    return amount


def _resolve_columns(headers: list[str], mapping: dict[str, str] | None) -> dict[str, int]:
    normalized = [normalize_header(header) for header in headers]
    if len(set(normalized)) != len(normalized):
        raise ImportFileError("Column headers must be unique. Rename any duplicate headers and try again.")

    columns: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        for index, header in enumerate(normalized):
            if header in aliases:
                columns[field] = index
                break

    if mapping:
        for field, selected_header in mapping.items():
            if field not in FIELDS:
                raise ImportFileError(f"Unsupported mapped field: {field}")
            try:
                columns[field] = headers.index(selected_header)
            except ValueError:
                raise ImportFileError(f"Mapped column for {FIELD_LABELS[field]} is not in this file") from None
    return columns


def _read_xlsx(content: bytes) -> tuple[list[str], list[tuple[int, tuple[Any, ...]]]]:
    if not content.startswith(b"PK") or not zipfile.is_zipfile(io.BytesIO(content)):
        raise ImportFileError("This is not a valid .xlsx workbook")

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_XLSX_ENTRIES:
                raise ImportFileError("The workbook contains too many internal files")
            if sum(info.file_size for info in infos) > MAX_UNCOMPRESSED_XLSX_BYTES:
                raise ImportFileError("The uncompressed workbook is too large")
            if any(info.flag_bits & 0x1 for info in infos):
                raise ImportFileError("Password-protected workbooks are not supported")
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise ImportFileError("The uploaded .xlsx file is not an Excel workbook")
            if any(
                name == "xl/vbaProject.bin" or name.startswith("xl/activeX/")
                for name in names
            ):
                raise ImportFileError("Macro-enabled workbooks are not supported")
    except zipfile.BadZipFile:
        raise ImportFileError("The Excel workbook is corrupted") from None

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if len(workbook.worksheets) != 1:
            workbook.close()
            raise ImportFileError("Use a workbook with a single worksheet")
        worksheet = workbook.active
        if worksheet.max_column > MAX_COLUMNS:
            workbook.close()
            raise ImportFileError(f"The worksheet cannot have more than {MAX_COLUMNS} columns")

        rows: list[tuple[int, tuple[Any, ...]]] = []
        headers: list[str] = []
        for row_number, row in enumerate(worksheet.iter_rows(values_only=False), start=1):
            if row_number > MAX_IMPORT_ROWS + 1:
                workbook.close()
                raise ImportFileError(f"The worksheet exceeds the {MAX_IMPORT_ROWS:,}-row import limit")
            if any(cell.data_type == "f" for cell in row):
                workbook.close()
                raise ImportFileError("Formula cells are not supported. Replace formulas with values and try again.")
            values = tuple(cell.value for cell in row)
            if row_number == 1:
                headers = [_string_value(value) for value in values]
            else:
                rows.append((row_number, values))
        workbook.close()
        return headers, rows
    except ImportFileError:
        raise
    except Exception as exc:
        raise ImportFileError("Could not read this Excel workbook. Check that it is not corrupted or password-protected.") from exc


def _read_csv(content: bytes) -> tuple[list[str], list[tuple[int, tuple[Any, ...]]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ImportFileError("CSV files must use UTF-8 text encoding") from None

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t") if sample else csv.excel
    except csv.Error:
        dialect = csv.excel

    try:
        reader = csv.reader(io.StringIO(text, newline=""), dialect)
        headers = next(reader, [])
        rows = []
        for row_number, row in enumerate(reader, start=2):
            if row_number > MAX_IMPORT_ROWS + 1:
                raise ImportFileError(f"The file exceeds the {MAX_IMPORT_ROWS:,}-row import limit")
            rows.append((row_number, tuple(row)))
        return [header.strip() for header in headers], rows
    except csv.Error:
        raise ImportFileError("The CSV file is malformed") from None


def _validated_values(raw_values: dict[str, Any], columns: dict[str, int]) -> tuple[dict[str, Any] | None, list[str]]:
    errors: list[str] = []
    missing = [field for field in REQUIRED_FIELDS if field not in columns]
    for field in missing:
        errors.append(f"Required column is not mapped: {FIELD_LABELS[field]}")

    def get(field: str) -> Any:
        index = columns.get(field)
        if index is None or index >= len(raw_values["cells"]):
            return None
        return raw_values["cells"][index]

    values: dict[str, Any] = {}
    if not missing:
        try:
            values["transaction_date"] = _parse_date(get("transaction_date"))
        except ValueError as exc:
            errors.append(str(exc))

        try:
            values["amount"] = _parse_amount(get("amount"))
        except ValueError as exc:
            errors.append(str(exc))

        raw_kind = _string_value(get("kind")).casefold()
        if raw_kind not in {"income", "expense"}:
            errors.append("Type must be either income or expense")
        else:
            values["kind"] = raw_kind

        raw_currency = (_string_value(get("currency")).upper() if "currency" in columns else "") or "USD"
        if raw_currency not in CURRENCIES:
            errors.append(f"Currency must be one of: {', '.join(sorted(CURRENCIES))}")
        else:
            values["currency"] = raw_currency

        category = _string_value(get("category"))
        if not category:
            errors.append("Category is required")
        elif len(category) > 80:
            errors.append("Category cannot exceed 80 characters")
        else:
            values["category"] = category

        merchant = _string_value(get("merchant")) if "merchant" in columns else ""
        note = _string_value(get("note")) if "note" in columns else ""
        if len(merchant) > 160:
            errors.append("Merchant cannot exceed 160 characters")
        if len(note) > 2000:
            errors.append("Note / Description cannot exceed 2,000 characters")
        values["merchant"] = merchant
        values["note"] = note

        raw_tags = _string_value(get("tags")) if "tags" in columns else ""
        values["tags"] = [tag.strip() for tag in raw_tags.split(",") if tag.strip()]

    if errors:
        return None, errors

    try:
        transaction = TransactionCreate.model_validate(values)
    except ValidationError as exc:
        return None, [error["msg"] for error in exc.errors()]

    return {
        "transaction_date": transaction.transaction_date.isoformat(),
        "kind": transaction.kind,
        "amount": str(transaction.amount),
        "currency": transaction.currency,
        "category": transaction.category.strip(),
        "merchant": transaction.merchant.strip(),
        "note": transaction.note.strip(),
        "tags": transaction.tags,
    }, []


def parse_transaction_file(
    filename: str,
    content: bytes,
    mapping: dict[str, str] | None = None,
) -> dict[str, Any]:
    extension = Path(filename or "").suffix.casefold()
    if extension not in {".xlsx", ".csv"}:
        raise ImportFileError("Choose an .xlsx or .csv file")
    if not content:
        raise ImportFileError("The selected file is empty")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ImportFileError(f"Files must be {MAX_UPLOAD_BYTES // (1024 * 1024)} MB or smaller")

    if extension == ".xlsx":
        headers, source_rows = _read_xlsx(content)
    else:
        headers, source_rows = _read_csv(content)

    if not headers or all(not header for header in headers):
        raise ImportFileError("The first row must contain column headers")
    if len(headers) > MAX_COLUMNS:
        raise ImportFileError(f"Files cannot have more than {MAX_COLUMNS} columns")

    columns = _resolve_columns(headers, mapping)
    unmapped_required = [
        field for field in REQUIRED_FIELDS
        if field not in columns
    ]
    unmapped_fields = [field for field in FIELDS if field not in columns]
    rows: list[dict[str, Any]] = []
    for row_number, cells in source_rows:
        if not any(value is not None and str(value).strip() for value in cells):
            continue
        if len(cells) > len(headers):
            parsed, errors = None, ["This row has more cells than the header row"]
        else:
            parsed, errors = _validated_values({"cells": cells}, columns)
        display_values = {
            FIELD_LABELS[field]: _string_value(cells[index]) if index < len(cells) else ""
            for field, index in columns.items()
        }
        rows.append({
            "row_number": row_number,
            "values": parsed or display_values,
            "valid": parsed is not None,
            "errors": errors,
            "possible_duplicate": False,
            "duplicate_reason": None,
        })

    valid_count = sum(row["valid"] for row in rows)
    return {
        "headers": headers,
        "mapped_columns": {field: headers[index] for field, index in columns.items()},
        "unmapped_fields": unmapped_fields,
        "unmapped_required": unmapped_required,
        "rows": rows,
        "summary": {
            "total_rows": len(rows),
            "valid_rows": valid_count,
            "invalid_rows": len(rows) - valid_count,
            "possible_duplicates": 0,
        },
    }


def transaction_signature(values: dict[str, Any]) -> tuple[Any, ...]:
    return (
        values["transaction_date"],
        values["kind"].casefold(),
        Decimal(str(values["amount"])).quantize(Decimal("0.01")),
        values["currency"].upper(),
        values["category"].strip().casefold(),
        values.get("merchant", "").strip().casefold(),
        values.get("note", "").strip().casefold(),
        tuple(sorted(tag.strip().casefold() for tag in values.get("tags", []))),
    )


def sample_template_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Transactions"
    worksheet.append(TEMPLATE_HEADERS)
    worksheet.append([date.today().isoformat(), "expense", 12.50, "USD", "Food", "Example merchant", "Optional note", "home, variable"])
    worksheet.freeze_panes = "A2"
    for column, width in enumerate((18, 12, 14, 12, 20, 24, 32, 24), start=1):
        worksheet.column_dimensions[chr(64 + column)].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()